"""Update the current Senior Geologist CV with a short Rocksurv showcase.
Keep the existing Word layout and mirrored DrawingML/VML text boxes.
"""
from copy import deepcopy
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from shutil import copy2
from zipfile import ZipFile, ZIP_DEFLATED
import json
import re
from docx import Document
from docx.shared import Mm
from lxml import etree
from PIL import Image, ImageDraw
from reportlab.graphics.barcode.qr import QrCodeWidget

ROOT = Path(__file__).resolve().parents[1]
REVIEW = ROOT / "review/cv-update-2026-10-09"
BACKUPS = ROOT / ".analysis/cv-update-2026-10-09/originals"
PORTFOLIO = "https://tntomg.github.io/en/"
QR_PATH = REVIEW / "rocksurv-portfolio-qr.png"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PIC = "http://schemas.openxmlformats.org/drawingml/2006/picture"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W, "r": R, "wp": WP, "a": A, "pic": PIC}
COPY = {
 "RU": {
  "subtitle": "Геология и разработка приложений Rocksurv | QA/QC",
  "profile": "Ведущий геолог с 15-летним опытом полевых работ, картирования, петрографии и геохимии. Разрабатываю приложения Rocksurv для ускорения сбора, проверки и обработки геологических данных. Объединяю геологическую практику, программирование и QA/QC.",
  "cta": "Посмотрите приложения и демонстрации на сайте",
  "apps_heading": "ПРИЛОЖЕНИЯ ROCKSURV",
  "apps": [
   ("Rocksurv Field", "полевые наблюдения, пробы и офлайн-карты."),
   ("Rocksurv Classifier", "классификация пород и проверка структурных замеров."),
   ("Rocksurv Core", "цифровая документация керна и горных выработок."),
   ("RockSurv Petrography", "определение минералов под микроскопом и атлас шлифов."),
  ],
  "status": "Разрабатываю прототипы. Сценарии работы и демонстрации доступны на сайте.",
  "education_heading": "Образование",
  "field_heading": "Опыт полевых работ",
 },
 "EN": {
  "subtitle": "Geology and Rocksurv App Development | QA/QC",
  "profile": "Senior Geologist with 15+ years of experience in fieldwork, mapping, petrography and geochemistry. I develop Rocksurv applications to accelerate geological data collection, validation and processing, combining geological practice, programming and QA/QC.",
  "cta": "Explore the applications and demonstrations online",
  "apps_heading": "ROCKSURV APPLICATIONS",
  "apps": [
   ("Rocksurv Field", "field observations, sampling and offline maps."),
   ("Rocksurv Classifier", "rock classification and structural measurement validation."),
   ("Rocksurv Core", "digital drill-core and mine-workings documentation."),
   ("RockSurv Petrography", "microscopic mineral identification and a thin-section atlas."),
  ],
  "status": "Prototypes in development. Workflows and demonstrations are available online.",
  "education_heading": "Education",
 },
}

def qn(namespace, local):
 return "{" + namespace + "}" + local

def paragraph_text(paragraph):
 return "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))

def make_run(text, rpr=None, bold=False, color=None):
 run = etree.Element(qn(W, "r"))
 props = deepcopy(rpr) if rpr is not None else etree.Element(qn(W, "rPr"))
 for tag in ["b", "bCs"]:
  old = props.find(qn(W, tag))
  if old is not None: props.remove(old)
 if bold: etree.SubElement(props, qn(W, "b"))
 if color:
  old = props.find(qn(W, "color"))
  if old is not None: props.remove(old)
  etree.SubElement(props, qn(W, "color")).set(qn(W, "val"), color)
 run.append(props)
 content = etree.SubElement(run, qn(W, "t"))
 content.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
 content.text = text
 return run

def replace_text(paragraph, text):
 runs = paragraph.xpath("./w:r[w:t]", namespaces=NS)
 props = deepcopy(runs[0].find(qn(W, "rPr"))) if runs else None
 for child in list(paragraph):
  if child.tag != qn(W, "pPr"): paragraph.remove(child)
 bold = props is not None and props.find(qn(W, "b")) is not None
 paragraph.append(make_run(text, props, bold=bold))

def new_paragraph(template, segments, heading=False):
 paragraph = deepcopy(template)
 for child in list(paragraph):
  if child.tag != qn(W, "pPr"): paragraph.remove(child)
 props = paragraph.find(qn(W, "pPr"))
 if props is None: props = etree.SubElement(paragraph, qn(W, "pPr"))
 spacing = props.find(qn(W, "spacing"))
 if spacing is None: spacing = etree.SubElement(props, qn(W, "spacing"))
 spacing.set(qn(W, "before"), "0")
 spacing.set(qn(W, "after"), "120" if heading else "100")
 for text, bold in segments:
  template_runs = template.xpath("./w:r[w:t]", namespaces=NS)
  rpr = template_runs[0].find(qn(W, "rPr")) if template_runs else None
  paragraph.append(make_run(text, rpr, bold=bold))
 return paragraph

def generate_qr():
 REVIEW.mkdir(parents=True, exist_ok=True)
 qr = QrCodeWidget(PORTFOLIO, barLevel="Q")
 qr.qr.make()
 modules, quiet, cell = qr.qr.getModuleCount(), 4, 16
 image = Image.new("RGB", ((modules + 2 * quiet) * cell,) * 2, "white")
 draw = ImageDraw.Draw(image)
 for row in range(modules):
  for col in range(modules):
   if qr.qr.isDark(row, col):
    x, y = (col + quiet) * cell, (row + quiet) * cell
    draw.rectangle((x, y, x + cell - 1, y + cell - 1), fill="black")
 image.save(QR_PATH, dpi=(500, 500))

def qr_drawing(image_rid, url_rid, next_id):
 doc = Document()
 picture = doc.add_paragraph().add_run().add_picture(str(QR_PATH), width=Mm(29))
 inline = deepcopy(picture._inline)
 anchor = etree.Element(qn(WP, "anchor"), nsmap={"wp": WP})
 anchor.attrib.update({
  "distT": "0", "distB": "0", "distL": "0", "distR": "0", "simplePos": "0",
  "relativeHeight": "251660000", "behindDoc": "0", "locked": "0",
  "layoutInCell": "1", "allowOverlap": "1",
 })
 etree.SubElement(anchor, qn(WP, "simplePos"), x="0", y="0")
 horizontal = etree.SubElement(anchor, qn(WP, "positionH"), relativeFrom="page")
 etree.SubElement(horizontal, qn(WP, "posOffset")).text = str(int(Mm(173)))
 vertical = etree.SubElement(anchor, qn(WP, "positionV"), relativeFrom="page")
 etree.SubElement(vertical, qn(WP, "posOffset")).text = str(int(Mm(2)))
 for child in list(inline):
  if child.tag == qn(WP, "extent"): anchor.append(child)
 etree.SubElement(anchor, qn(WP, "wrapNone"))
 for child in list(inline): anchor.append(child)
 docpr = anchor.find(qn(WP, "docPr"))
 docpr.set("id", str(next_id))
 docpr.set("name", "Rocksurv portfolio QR")
 docpr.set("descr", "Rocksurv applications and demonstrations: " + PORTFOLIO)
 etree.SubElement(docpr, qn(A, "hlinkClick")).set(qn(R, "id"), url_rid)
 anchor.xpath(".//a:blip", namespaces=NS)[0].set(qn(R, "embed"), image_rid)
 for element in anchor.xpath(".//pic:cNvPr", namespaces=NS):
  element.set("id", str(next_id))
  etree.SubElement(element, qn(A, "hlinkClick")).set(qn(R, "id"), url_rid)
 drawing = etree.Element(qn(W, "drawing")); drawing.append(anchor)
 run = etree.Element(qn(W, "r")); run.append(drawing)
 return run

def update_language(lang):
 source = ROOT / "public/cv" / f"Aleksei-Pakhalko-CV-{lang}-2026.docx"
 BACKUPS.mkdir(parents=True, exist_ok=True)
 backup = BACKUPS / source.name
 if not backup.exists(): copy2(source, backup)
 for filename in [source.name.replace(".docx", ".pdf"), source.name.replace("Aleksei", "Alexey"), source.name.replace("Aleksei", "Alexey").replace(".docx", ".pdf")]:
  existing = source.parent / filename
  if existing.exists() and not (BACKUPS / filename).exists(): copy2(existing, BACKUPS / filename)
 with ZipFile(backup) as archive:
  files = {item.filename: archive.read(item.filename) for item in archive.infolist()}
 tree = etree.fromstring(files["word/document.xml"])
 rels = etree.fromstring(files["word/_rels/document.xml.rels"])
 existing_ids = {rel.get("Id") for rel in rels}
 next_rel = max([int(rid[3:]) for rid in existing_ids if rid.startswith("rId") and rid[3:].isdigit()] + [0]) + 1
 image_rid, url_rid = "rId" + str(next_rel), "rId" + str(next_rel + 1)
 etree.SubElement(rels, qn(REL, "Relationship"), Id=image_rid, Type=R + "/image", Target="media/rocksurv-portfolio-qr.png")
 etree.SubElement(rels, qn(REL, "Relationship"), Id=url_rid, Type=R + "/hyperlink", Target=PORTFOLIO, TargetMode="External")
 body = tree.find(qn(W, "body")); body_paras = body.findall(qn(W, "p")); texts = COPY[lang]
 replace_text(body_paras[3], texts["subtitle"])
 replace_text(body_paras[4], texts["profile"])
 p = body_paras[4]; style = deepcopy(p.xpath("./w:r/w:rPr", namespaces=NS)[0])
 spacing = p.find(qn(W, "pPr")).find(qn(W, "spacing"))
 if spacing is None: spacing = etree.SubElement(p.find(qn(W, "pPr")), qn(W, "spacing"))
 spacing.set(qn(W, "after"), "500" if lang == "RU" else "360")
 line = etree.SubElement(p, qn(W, "r")); etree.SubElement(line, qn(W, "br"))
 p.append(make_run(texts["cta"] + ": ", style, bold=True))
 link = etree.SubElement(p, qn(W, "hyperlink")); link.set(qn(R, "id"), url_rid)
 url_run = make_run("tntomg.github.io/en/", style, color="2E6B34")
 etree.SubElement(url_run.find(qn(W, "rPr")), qn(W, "u")).set(qn(W, "val"), "single")
 link.append(url_run)
 logo_runs = [r for r in body_paras[0].findall(qn(W, "r")) if r.xpath(".//a:blip", namespaces=NS)]
 assert len(logo_runs) == 4, (lang, "Unexpected header image count", len(logo_runs))
 for run in logo_runs[1:]: body_paras[0].remove(run)
 max_id = max([int(e.get("id")) for e in tree.xpath("//wp:docPr", namespaces=NS)] + [0])
 body_paras[0].append(qr_drawing(image_rid, url_rid, max_id + 1))
 boxes = tree.xpath("//w:txbxContent", namespaces=NS)
 for index in (4, 5):
  replace_text(boxes[index][0], "+7 904 637-44-16 / +7 927 234-45-25")
 if lang == "EN":
  city_anchor = boxes[0].xpath("ancestor::wp:anchor", namespaces=NS)[0]
  x = city_anchor.find(qn(WP, "positionH")).find(qn(WP, "posOffset"))
  y = city_anchor.find(qn(WP, "positionV")).find(qn(WP, "posOffset"))
  x.text = str(int(x.text) + int(Mm(5)))
  y.text = str(int(y.text) - 63500)
  legacy = boxes[1].getparent().getparent()
  style = legacy.get("style")
  style = re.sub(r"margin-left:([0-9.]+)pt", lambda m: "margin-left:" + str(float(m[1]) + 14.173228) + "pt", style)
  style = re.sub(r"margin-top:([0-9.]+)pt", lambda m: "margin-top:" + str(float(m[1]) - 5) + "pt", style)
  legacy.set("style", style)

 for first, second in [(6, 12), (7, 13)]:
  sidebar = boxes[first]; records = sidebar.findall(qn(W, "p"))
  education_at = next(i for i, row in enumerate(records) if paragraph_text(row).strip() == texts["education_heading"])
  education = [deepcopy(row) for row in records[education_at:] if paragraph_text(row).strip()]
  for row in records[education_at:]: sidebar.remove(row)
  heading_template, body_template = records[0], records[1]
  apps = [new_paragraph(heading_template, [(texts["apps_heading"], True)], heading=True)]
  for title, summary in texts["apps"]:
   apps.append(new_paragraph(body_template, [(title, True), (": " + summary, False)]))
  apps.append(new_paragraph(body_template, [(texts["status"], False)]))
  for row in reversed(apps): sidebar.insert(0, row)
  certificates = boxes[second]
  while len(certificates) and not paragraph_text(certificates[-1]).strip(): certificates.remove(certificates[-1])
  for row in education: certificates.append(row)
 if lang == "RU":
  for index in (10, 11): replace_text(boxes[index][0], texts["field_heading"])
 metadata = etree.fromstring(files["docProps/core.xml"])
 modified = metadata.find("{http://purl.org/dc/terms/}modified")
 if modified is not None: modified.text = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
 files["docProps/core.xml"] = etree.tostring(metadata, encoding="UTF-8", xml_declaration=True, standalone=True)
 files["word/document.xml"] = etree.tostring(tree, encoding="UTF-8", xml_declaration=True, standalone=True)
 files["word/_rels/document.xml.rels"] = etree.tostring(rels, encoding="UTF-8", xml_declaration=True, standalone=True)
 files["word/media/rocksurv-portfolio-qr.png"] = QR_PATH.read_bytes()
 content_types = etree.fromstring(files["[Content_Types].xml"])
 if not any(e.get("Extension") == "png" for e in content_types):
  etree.SubElement(content_types, "{http://schemas.openxmlformats.org/package/2006/content-types}Default", Extension="png", ContentType="image/png")
 files["[Content_Types].xml"] = etree.tostring(content_types, encoding="UTF-8", xml_declaration=True, standalone=True)
 updated = ROOT / "CV" / f"CV_Pakhalko_Senior_Geologist_Rocksurv_2026_{lang}.docx"
 with ZipFile(updated, "w", compression=ZIP_DEFLATED) as archive:
  for filename, payload in files.items(): archive.writestr(filename, payload)
 copy2(updated, source)
 copy2(updated, source.with_name(source.name.replace("Aleksei", "Alexey")))
 print(f"Updated {lang}: source + public canonical and alias; original retained")
 return {"language": lang, "source": str(updated), "publicDocx": str(source), "sourceSha256": sha256(backup.read_bytes()).hexdigest(), "updatedSha256": sha256(updated.read_bytes()).hexdigest(), "portfolioUrl": PORTFOLIO}

if __name__ == "__main__":
 generate_qr()
 changes = [update_language(lang) for lang in ("RU", "EN")]
 (REVIEW / "update-manifest.json").write_text(json.dumps(changes, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
